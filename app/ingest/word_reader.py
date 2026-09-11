"""Word 资料来源 — 实现 SourceReader 接口。

企业最常见数据格式，支持 .docx 文件按标题分块。

分块策略：
1. 优先按 Word 标题样式（Heading 1/2/3...）分块（推荐用法）
2. 如果没检测到任何标题样式，自动 fallback 到中文编号标题启发式检测：
   - 一、二、三... 一级标题
   - （一）（二）... 二级标题
   - 1. 2. 3.... 三级标题
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import List, Tuple

from app.contracts import Chunk, ChunkMeta, ChunkType, ChunkTypeClassification, SourceReader
from app.cross.logging import get_logger

logger = get_logger(__name__)

# 标题样式匹配（大小写不敏感，兼容不同版本 python-docx 和国际化）
_HEADING_STYLE_RE = re.compile(r"^(heading|标题)\s*(\d+)$", re.IGNORECASE)
_TITLE_STYLE_RE = re.compile(r"^(title|标题)$", re.IGNORECASE)

# 中文编号标题启发式匹配（用于 fallback 模式，当没有样式标记时）
# 一级：一、 二、 三、 ...
# 二级：（一）（二）... 或者 (一) (二) ...
# 三级：1. 2. 3. ... 或者 1、 2、 3、 ...
# 四级：(1) (2) ... 或者 （1）（2）...
_NUM_TITLE_RE = [
    # (pattern, level)
    (re.compile(r"^([一二三四五六七八九十]+)、"), 1),  # 一、二、
    (re.compile(r"^（([一二三四五六七八九十]+)）"), 2),  # （一）
    (re.compile(r"^\(([一二三四五六七八九十]+)\)"), 2),  # (一)
    (re.compile(r"^(\d+)\."), 3),  # 1.
    (re.compile(r"^(\d+)、"), 3),  # 1、
    (re.compile(r"^（(\d+)）"), 4),  # （1）
    (re.compile(r"^\((\d+)\)"), 4),  # (1)
]


class WordReader(SourceReader):
    """Word 文档读取，按标题分块。

    每个标题 + 正文 = 一个 Chunk。
    标题层级信息保留在 ChunkMeta.heading_level 中。

    分块策略：
    - 优先：按 Word 标题样式（Heading 1/2/3）分块
    - Fallback：如果没有检测到任何标题样式，自动启发式检测中文编号标题（一、二、三...）

    用法::
        reader = WordReader()
        chunks = reader.read("文件.docx")
    """

    def read(self, source: str, doc_name: str = "") -> List[Chunk]:
        path = Path(source)
        if not doc_name:
            doc_name = path.stem
        doc_id = hashlib.md5(str(path).encode()).hexdigest()[:16]

        try:
            from docx import Document
            doc = Document(str(path))
        except ImportError:
            raise ImportError("python-docx is required for WordReader: pip install python-docx")
        except Exception as e:
            logger.error("failed to open Word document: %s", e)
            return []

        logger.info("reading Word document: %s", path)

        # 先尝试：按样式分块
        chunks_by_style = self._split_by_style(doc, doc_id, doc_name)

        # 如果只分出一个大 chunk，说明文档没有用标题样式，fallback 到启发式编号检测
        if len(chunks_by_style) <= 1:
            chunks = self._split_by_heuristic(doc, doc_id, doc_name)
            # 启发式会"跳过首个编号标题之前的内容"（关键词海绵）。若文档根本没有编号
            # 标题，启发式会把整篇丢光（返回 0）。此时必须保留按样式得到的那一个块，
            # 否则整篇文档对检索完全不可见。
            if chunks:
                logger.info("WordReader: fallback to heuristic splitting, %d chunks from %s", len(chunks), path)
                return chunks
            logger.info("WordReader: heuristic gave 0 chunks, keep style-based result from %s", path)

        logger.info("WordReader: %d chunks from %s (style-based)", len(chunks_by_style), path)
        return chunks_by_style

    def _split_by_style(self, doc, doc_id, doc_name) -> List[Chunk]:
        """按 Word 标题样式分块。"""
        chunks: List[Chunk] = []
        current_heading = ""
        current_level = 0
        current_paragraphs: List[str] = []
        chunk_idx = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # 判断是否是标题（大小写不敏感 + 正则回退）
            style_name = para.style.name if para.style else ""
            level = 0
            is_heading = False

            m = _TITLE_STYLE_RE.match(style_name)
            if m:
                is_heading = True
                level = 0
            else:
                m = _HEADING_STYLE_RE.match(style_name)
                if m:
                    is_heading = True
                    level = int(m.group(2))

            if is_heading:
                # 保存上一个标题的内容
                if current_paragraphs:
                    chunks.append(self._make_chunk(
                        content="\n".join(current_paragraphs),
                        doc_id=doc_id,
                        doc_name=doc_name,
                        heading_number=current_heading,
                        heading_level=current_level,
                        chunk_index=chunk_idx,
                    ))
                    chunk_idx += 1
                current_heading = text
                current_level = level
                current_paragraphs = []
            else:
                current_paragraphs.append(text)

        # 最后一个标题
        if current_paragraphs:
            chunks.append(self._make_chunk(
                content="\n".join(current_paragraphs),
                doc_id=doc_id,
                doc_name=doc_name,
                heading_number=current_heading,
                heading_level=current_level,
                chunk_index=chunk_idx,
            ))

        return chunks

    def _split_by_heuristic(self, doc, doc_id, doc_name) -> List[Chunk]:
        """启发式分块：检测中文编号标题（一、二、三...）。

        跳过第一个编号标题之前的所有内容（通常是文档标题、发文字号、收文单位等），
        避免这些通用信息成为关键词海绵，吸走所有查询。
        """
        chunks: List[Chunk] = []
        current_heading = ""
        current_level = 0
        current_lines: List[str] = []
        found_first_heading = False
        chunk_idx = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # 检查本行是否是编号标题
            is_heading, level = self._is_number_heading(text)

            # 第一个标题之前的内容直接跳过（关键词海绵问题）
            if not found_first_heading:
                if is_heading:
                    found_first_heading = True
                    current_heading = text
                    current_level = level
                    current_lines = []
                continue
            else:
                if is_heading:
                    # 保存上一个标题的内容
                    if current_lines:
                        chunks.append(self._make_chunk(
                            content="\n".join(current_lines),
                            doc_id=doc_id,
                            doc_name=doc_name,
                            heading_number=current_heading,
                            heading_level=current_level,
                            chunk_index=chunk_idx,
                        ))
                        chunk_idx += 1
                    current_heading = text
                    current_level = level
                    current_lines = []
                else:
                    current_lines.append(text)

        # 最后一个标题
        if found_first_heading and current_lines:
            chunks.append(self._make_chunk(
                content="\n".join(current_lines),
                doc_id=doc_id,
                doc_name=doc_name,
                heading_number=current_heading,
                heading_level=current_level,
                chunk_index=chunk_idx,
            ))

        return chunks

    def _is_number_heading(self, text: str) -> Tuple[bool, int]:
        """检查文本行是否是编号标题，返回 (是否是标题, 层级)。"""
        for pattern, level in _NUM_TITLE_RE:
            if pattern.match(text):
                return True, level
        return False, 0

    def _make_chunk(
        self,
        content: str,
        doc_id: str,
        doc_name: str,
        heading_number: str,
        heading_level: int,
        chunk_index: int = 0,
    ) -> Chunk:
        if heading_number:
            content = heading_number + "\n" + content
        meta = ChunkMeta(
            doc_id=doc_id,
            doc_name=doc_name,
            mcu_model="",
            page_num=0,
            heading_number=heading_number,
            heading_title=heading_number,
            heading_level=heading_level,
            chunk_type=ChunkType.TEXT,
            chunk_index=chunk_index,
            token_count=len(content) // 4,
            classification=ChunkTypeClassification.PUBLIC,
        )
        return Chunk(
            content=content.strip(),
            meta=meta,
        )
