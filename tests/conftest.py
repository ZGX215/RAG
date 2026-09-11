"""测试共享夹具 — 多来源数据接口。"""

from __future__ import annotations

import sys
from pathlib import Path

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

import os
import tempfile

import pytest


@pytest.fixture
def sample_markdown():
    """创建 Markdown 测试临时文件，返回 (file_path, expected_chunks)。"""
    content = """# 第一章 总则

第一条 为了保护个人信息权益，规范个人信息处理活动，促进个人信息合理利用，根据宪法，制定本法。

第二条 自然人的个人信息受法律保护，任何组织、个人不得侵害自然人的个人信息权益。

## 第二节 基本原则

第四条 个人信息是以电子或者其他方式记录的与已识别或者可识别的自然人有关的各种信息。

### 1. 核心概念

个人信息包括姓名、出生日期、身份证件号码、生物识别信息、住址、电话号码、电子邮箱、健康信息、行踪信息等。
"""
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".md", encoding="utf-8", delete=False
    ) as f:
        f.write(content)
        file_path = f.name
    yield file_path, 3  # 3 个标题
    os.unlink(file_path)


@pytest.fixture
def sample_md_no_headings():
    """无标题的 Markdown 文件。"""
    content = "第一段纯文本内容，没有标题。\n\n第二段内容。\n\n第三段内容。"
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".md", encoding="utf-8", delete=False
    ) as f:
        f.write(content)
        file_path = f.name
    yield file_path, 1  # 无标题 → 兜底把全文作为 1 个块（不能丢数据）
    os.unlink(file_path)


@pytest.fixture
def sample_md_empty():
    """空 Markdown 文件。"""
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".md", encoding="utf-8", delete=False
    ) as f:
        f.write("")
        file_path = f.name
    yield file_path, 0
    os.unlink(file_path)


@pytest.fixture
def sample_txt_articles():
    """法律条款文本文件。"""
    content = """第一条 为了保护个人信息权益，规范个人信息处理活动，促进个人信息合理利用，根据宪法，制定本法。

第二条 自然人的个人信息受法律保护，任何组织、个人不得侵害自然人的个人信息权益。

第三条 本法所称个人信息，是指以电子或者其他方式记录的与已识别或者可识别的自然人有关的各种信息，不包括匿名化处理后的信息。
"""
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", encoding="utf-8", delete=False
    ) as f:
        f.write(content)
        file_path = f.name
    yield file_path, 3  # 3 个条款
    os.unlink(file_path)


@pytest.fixture
def sample_txt_paragraphs():
    """普通段落文本文件。"""
    content = "第一段内容：这里是一段比较长的文本，用于测试段落分块功能。\n\n第二段内容：这里也是一段较长的文本，确保不会被合并。\n\n第三段内容：继续测试段落分块，这段也是够长的。\n\n第四段内容：最后一段测试文本，确保四段分别独立。"
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", encoding="utf-8", delete=False
    ) as f:
        f.write(content)
        file_path = f.name
    yield file_path, 4  # 4 个段落
    os.unlink(file_path)


@pytest.fixture
def sample_txt_empty():
    """空文本文件。"""
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", encoding="utf-8", delete=False
    ) as f:
        f.write("")
        file_path = f.name
    yield file_path, 0
    os.unlink(file_path)


@pytest.fixture
def sample_txt_mixed_content():
    """不含条款号标记的普通文本（确认不分条款模式）。"""
    content = "个人信息保护法是一部重要的法律，于2021年8月20日通过，2021年11月1日正式施行。\n\n该法主要规范个人信息的处理活动，保护自然人个人信息权益。\n\n任何组织和个人都应当遵守个人信息保护法的相关规定，不得非法收集、使用个人信息。"
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", encoding="utf-8", delete=False
    ) as f:
        f.write(content)
        file_path = f.name
    yield file_path, 3
    os.unlink(file_path)


@pytest.fixture
def sample_docx():
    """创建 Word 测试文档文件，返回 (file_path, expected_chunks)。"""
    from docx import Document
    doc = Document()
    doc.add_heading("第一章 总则", level=1)
    doc.add_paragraph("第一条 为了保护个人信息权益，规范个人信息处理活动，促进个人信息合理利用，根据宪法，制定本法。")
    doc.add_paragraph("第二条 自然人的个人信息受法律保护。")
    doc.add_heading("第二节 基本原则", level=2)
    doc.add_paragraph("第四条 个人信息是以电子或者其他方式记录的。")
    doc.add_heading("1. 核心概念", level=3)
    doc.add_paragraph("个人信息包括姓名、出生日期、身份证件号码等。")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        file_path = f.name
    yield file_path, 3  # 3 个标题
    os.unlink(file_path)


@pytest.fixture
def sample_docx_no_headings():
    """无标题的 Word 文档。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("第一段文字内容。")
    doc.add_paragraph("第二段文字内容。")
    doc.add_paragraph("第三段文字内容。")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        doc.save(f.name)
        file_path = f.name
    yield file_path, 1  # 无标题 → 保留样式分块结果（整篇 1 个块），不能返回 0
    os.unlink(file_path)


# 跨平台候选 CJK 字体：Windows 用微软雅黑，Linux 用 Noto CJK，macOS 用苹方。
# 原先只写死了 Windows 路径，Linux CI 上会退化成 Helvetica，
# 渲染中文直接抛 FPDFUnicodeEncodingException（首次 CI 才暴露的问题）。
_CJK_FONT_CANDIDATES = (
    "C:/Windows/Fonts/msyh.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
    "/System/Library/Fonts/PingFang.ttc",
)


def _find_cjk_font() -> str | None:
    """查找可用的 CJK 字体，找不到返回 None。"""
    for path in _CJK_FONT_CANDIDATES:
        if os.path.exists(path):
            return path
    return None


def _create_pdf(file_path: str, pages: list[str]):
    """用 fpdf2 生成临时 PDF 文件。

    含中文时必须有 CJK 字体。找不到就 skip 而不是 error —— 让"缺字体"这种
    环境问题表现为"跳过"，而不是伪装成测试失败。
    """
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()

    has_chinese = any("\u4e00" <= c <= "\u9fff" for page_text in pages for c in page_text)
    if has_chinese:
        font_path = _find_cjk_font()
        if not font_path:
            pytest.skip(
                "缺少 CJK 字体，无法生成中文测试 PDF"
                "（Linux 可 apt install fonts-noto-cjk；Windows 自带微软雅黑）"
            )
        pdf.add_font("CJK", "", font_path)
        pdf.set_font("CJK", size=12)
    else:
        pdf.set_font("Helvetica", size=12)

    for i, page_text in enumerate(pages):
        if i > 0:
            pdf.add_page()
        pdf.multi_cell(0, 10, page_text)
    pdf.output(file_path)


@pytest.fixture
def sample_pdf_articles():
    """含法律条款的 PDF（应触发条款模式）。"""
    pages = [
        "第一条 为了保护个人信息权益，规范个人信息处理活动，促进个人信息合理利用，根据宪法，制定本法。",
        "第二条 自然人的个人信息受法律保护，任何组织、个人不得侵害自然人的个人信息权益。",
    ]
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        _create_pdf(f.name, pages)
        file_path = f.name
    yield file_path, 2  # 2 个条款
    os.unlink(file_path)


@pytest.fixture
def sample_pdf_normal():
    """普通 PDF（无法律条款，应触发逐页模式）。"""
    pages = [
        "This is page 1 of a normal document. It contains general information about the project.",
        "This is page 2 of the document. It continues with more details about the implementation.",
        "This is page 3 of the document. It concludes with final remarks and next steps.",
    ]
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        _create_pdf(f.name, pages)
        file_path = f.name
    yield file_path, 3  # 3 页
    os.unlink(file_path)


@pytest.fixture
def sample_pdf_single_page():
    """单页普通 PDF。"""
    pages = ["A single page document with some content about testing."]
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        _create_pdf(f.name, pages)
        file_path = f.name
    yield file_path, 1  # 1 页
    os.unlink(file_path)


@pytest.fixture
def sample_pdf_empty():
    """空 PDF（无内容页）。"""
    from fpdf import FPDF
    pdf = FPDF()
    # 空 PDF：有页面但无文本内容
    pdf.add_page()
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        pdf.output(f.name)
        file_path = f.name
    yield file_path, 0  # 0 个分块（页面无提取文本）
    os.unlink(file_path)